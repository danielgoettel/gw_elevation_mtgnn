import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

out_dir = Path(r"C:\Users\dgoet\DeepLearning\gw_elevation_mtgnn\DMG_Documentation\Program Flow")
out_dir.mkdir(parents=True, exist_ok=True)

# ── 1. Create the flow diagram ──────────────────────────────────────────

fig, ax = plt.subplots(figsize=(16, 22))
ax.set_xlim(0, 16)
ax.set_ylim(0, 22)
ax.axis('off')

colors = {
    'entry':  '#2C3E50',
    'config': '#8E44AD',
    'data':   '#2980B9',
    'graph':  '#27AE60',
    'model':  '#E67E22',
    'train':  '#C0392B',
    'eval':   '#16A085',
}

def box(ax, x, y, w, h, text, color, fontsize=9, bold=False):
    rect = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.15",
                          facecolor=color, edgecolor='white', linewidth=1.5, alpha=0.92)
    ax.add_patch(rect)
    weight = 'bold' if bold else 'normal'
    ax.text(x + w/2, y + h/2, text, ha='center', va='center',
            fontsize=fontsize, color='white', weight=weight, wrap=True,
            fontfamily='monospace')

def arrow(ax, x1, y1, x2, y2):
    ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle='->', color='#555555', lw=1.5))

def sub_box(ax, x, y, w, h, text, fontsize=7.5):
    rect = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.1",
                          facecolor='#ECF0F1', edgecolor='#95A5A6', linewidth=1, alpha=0.9)
    ax.add_patch(rect)
    ax.text(x + w/2, y + h/2, text, ha='center', va='center',
            fontsize=fontsize, color='#2C3E50', fontfamily='monospace')

# Title
ax.text(8, 21.5, 'gw_elevation_mtgnn \u2014 Program Flow', ha='center', va='center',
        fontsize=16, weight='bold', color='#2C3E50', fontfamily='sans-serif')

# ENTRY
box(ax, 5.5, 20.3, 5, 0.7, 'train_model.py: main()', colors['entry'], 11, bold=True)
arrow(ax, 8, 20.3, 8, 19.9)

# CONFIG
box(ax, 5, 19.1, 6, 0.7, 'train_config.py: define_base_configuration()\n-> config dict + generate_configurations()', colors['config'], 8.5)
arrow(ax, 8, 19.1, 8, 18.7)

# FOR EACH CONFIG border
rect = FancyBboxPatch((1, 1.5), 14, 17, boxstyle="round,pad=0.2",
                      facecolor='none', edgecolor='#7F8C8D', linewidth=2, linestyle='--')
ax.add_patch(rect)
ax.text(8, 18.3, 'FOR EACH CONFIG -> run_training_and_evaluation(config)', ha='center',
        fontsize=10, weight='bold', color='#7F8C8D', fontfamily='sans-serif')

# DATA PIPELINE (left column)
box(ax, 1.5, 16.5, 5.5, 0.7, 'process_data.py: main()', colors['data'], 10, bold=True)

sub_box(ax, 1.7, 15.4, 5.1, 0.45, 'preprocessing.py: piezometer_measurements()\nRead DINOloket CSVs -> filter "reliable"')
sub_box(ax, 1.7, 14.8, 5.1, 0.45, 'process_series()\nGroup by date, resample weekly -> complete_daily')
sub_box(ax, 1.7, 14.2, 5.1, 0.45, 'fill_and_select_data()\nselect_nodes() top-200 -> interpolate -> df_piezo')
sub_box(ax, 1.7, 13.6, 5.1, 0.45, 'load_external_data()\n4 pump + 2 prec + 2 evap + ~4 river -> weekly')
sub_box(ax, 1.7, 13.0, 5.1, 0.45, 'split_and_normalize_data()\nConcat [200|4|2|2|4]=212 -> 80/20 -> MinMaxScaler')

arrow(ax, 4.25, 16.5, 4.25, 15.9)
arrow(ax, 4.25, 15.4, 4.25, 15.3)
arrow(ax, 4.25, 14.8, 4.25, 14.7)
arrow(ax, 4.25, 14.2, 4.25, 14.1)
arrow(ax, 4.25, 13.6, 4.25, 13.5)

sub_box(ax, 1.7, 12.3, 5.1, 0.5, 'OUTPUT:\ntrain/val/test_data, masks, scaler\n[T x 212] DataFrames', 7)
arrow(ax, 4.25, 13.0, 4.25, 12.85)

# GRAPH CONSTRUCTION (right column)
box(ax, 9, 16.5, 5.5, 0.7, 'gnn_data_prep.py: main()', colors['graph'], 10, bold=True)

sub_box(ax, 9.2, 15.4, 5.1, 0.45, 'load_and_concatenate_metadata()\nPiezo+Pump+Prec+Evap+River coords -> 212 nodes')
sub_box(ax, 9.2, 14.8, 5.1, 0.45, 'create_static_features()\nNormalize x,y,z + type -> [212x4] tensor')
sub_box(ax, 9.2, 14.0, 5.1, 0.6, 'Generate adjacency matrix (by graph_type):\n"default" | "geolayer" | "rf fixed" | "rf variable"\nPiezo<->Piezo + Piezo->Exogenous connections')

arrow(ax, 11.75, 16.5, 11.75, 15.9)
arrow(ax, 11.75, 15.4, 11.75, 15.3)
arrow(ax, 11.75, 14.8, 11.75, 14.65)

sub_box(ax, 9.2, 13.2, 5.1, 0.5, 'OUTPUT:\nA_tilde [212x212], static_features [212x4]', 7.5)
arrow(ax, 11.75, 14.0, 11.75, 13.75)

# MODEL CREATION
arrow(ax, 4.25, 12.3, 8, 11.7)
arrow(ax, 11.75, 13.2, 8, 11.7)

box(ax, 5, 11, 6, 0.6, 'create_mtgnn_model() -> MTGNN or MTGNN_LSTM', colors['model'], 9, bold=True)

sub_box(ax, 2, 10.0, 12, 0.7, 'MTGNN Architecture:\nStart Conv -> 4x [DilatedInception -> MixProp(A) -> LayerNorm -> Skip] -> End Conv\nInput: [B, 1, 212, W+1]  ->  Output: [B, 1, 200, 1]', 8)
arrow(ax, 8, 11, 8, 10.75)

# TRAINING
arrow(ax, 8, 10.0, 8, 9.6)

box(ax, 3, 8.7, 10, 0.8, 'train() -- Progressive Future Windows\nFOR F_w = 1, 2, 3: warm-start from previous', colors['train'], 9, bold=True)

sub_box(ax, 1.5, 7.5, 6, 0.9, 'AutoregressiveTimeSeriesDataset:\ninput  = data[idx:idx+W, :200]     (piezo)\nforces = data[idx:idx+W+F, 200:]  (exo)\ntarget = data[idx+W:idx+W+F, :200]', 7.5)

sub_box(ax, 8.5, 7.5, 6, 0.9, 'Training Loop (per epoch):\nFOR each batch, FOR t=0..F_w:\n  Slice forces -> combine -> model forward\n  Shift window: drop oldest, append pred\nMasked MSE -> backprop -> optimizer', 7.5)

arrow(ax, 8, 8.7, 4.5, 8.45)
arrow(ax, 8, 8.7, 11.5, 8.45)

sub_box(ax, 3, 6.5, 10, 0.6, 'Early Stopping + ReduceLROnPlateau + Save best model per window', 8)
arrow(ax, 8, 7.5, 8, 7.15)

# EVALUATION
arrow(ax, 8, 6.5, 8, 6.1)

box(ax, 3, 5.3, 10, 0.7, 'EVALUATION', colors['eval'], 10, bold=True)

sub_box(ax, 1.5, 4.3, 4.2, 0.7, 'make_predictions()\n100-step autoregressive\ninverse_transform()', 7.5)
sub_box(ax, 6, 4.3, 4, 0.7, 'calculate_rmse\n_per_piezometer()\nSummarize by geolayer', 7.5)
sub_box(ax, 10.3, 4.3, 4.2, 0.7, 'plot_comparison_sequence()\nplot_rmse_3d_network()\nSave scatter HTML', 7.5)

arrow(ax, 5, 5.3, 3.6, 5.05)
arrow(ax, 8, 5.3, 8, 5.05)
arrow(ax, 11, 5.3, 12.4, 5.05)

# OUTPUT
arrow(ax, 8, 4.3, 8, 3.8)

box(ax, 3, 2.8, 10, 0.7, 'OUTPUT: summary CSV + RMSE JSON\nanalyze_results() -> best config', colors['entry'], 9, bold=True)

# Legend
legend_items = [
    ('Entry/Output', colors['entry']),
    ('Configuration', colors['config']),
    ('Data Pipeline', colors['data']),
    ('Graph Construction', colors['graph']),
    ('Model', colors['model']),
    ('Training', colors['train']),
    ('Evaluation', colors['eval']),
]
for i, (label, color) in enumerate(legend_items):
    rect = FancyBboxPatch((0.3, 1.8 - i*0.25), 0.4, 0.18, boxstyle="round,pad=0.02",
                          facecolor=color, edgecolor='white', alpha=0.9)
    ax.add_patch(rect)
    ax.text(0.85, 1.89 - i*0.25, label, fontsize=7.5, va='center', color='#2C3E50')

plt.tight_layout()
diagram_path = out_dir / 'program_flow_diagram.png'
fig.savefig(str(diagram_path), dpi=200, bbox_inches='tight', facecolor='white')
plt.close()
print(f"Diagram saved: {diagram_path}")

# ── 2. Create the Word document ─────────────────────────────────────────

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

style_h1 = doc.styles['Heading 1']
style_h1.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
style_h2 = doc.styles['Heading 2']
style_h2.font.color.rgb = RGBColor(0x29, 0x80, 0xB9)

title = doc.add_heading('gw_elevation_mtgnn \u2014 Program Flow Documentation', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Generated: February 2026')

doc.add_heading('Flow Diagram', level=1)
doc.add_picture(str(diagram_path), width=Inches(7))

doc.add_heading('Detailed Program Flow', level=1)

# 1. Entry
doc.add_heading('1. Entry Point \u2014 train_model.py: main()', level=2)
doc.add_paragraph(
    'The program starts in train_model.py main(). It loads the base configuration from '
    'train_config.py: define_base_configuration(), generates configuration variants via '
    'generate_configurations(), and iterates over each config calling '
    'run_training_and_evaluation(config).'
)

# 2. Data Pipeline
doc.add_heading('2. Data Pipeline \u2014 process_data.py: main()', level=2)
doc.add_paragraph(
    'Called first inside run_training_and_evaluation(). Responsible for loading, '
    'cleaning, and preparing all time series data.'
)

steps = [
    ('2a. Load Piezometer Data',
     'preprocessing.py: piezometer_measurements() reads DINOloket CSV files per well. '
     'Filters to "reliable" measurements only (BIJZONDERHEID column).'),
    ('2b. Process Series',
     'process_series() groups measurements by date, aggregates duplicates, and resamples '
     'to weekly frequency. Filters to post-2004 data. Returns complete_daily DataFrame.'),
    ('2c. Select Nodes & Fill Missing Data',
     'fill_and_select_data() calls select_nodes() to pick the top 200 piezometers with '
     'the least missing data. Creates a missing_data_mask (~isna), then interpolates '
     '(linear) and backfills remaining gaps. Returns df_piezo [T x 200].'),
    ('2d. Load External Data',
     'load_external_data() loads 4 data sources: pumping wells (4 cols), precipitation '
     '(2 cols), evaporation (2 cols), and river stages (~4 cols after filtering). All '
     'are resampled to weekly and aligned to the piezometer date range.'),
    ('2e. Split & Normalize',
     'split_and_normalize_data() concatenates piezometer data with external data into a '
     'single DataFrame [T x ~212]. Splits 80/20 into train/test (no shuffle, time-ordered). '
     'Val = Test. Fits MinMaxScaler on train, transforms all splits. Returns train_data, '
     'val_data, test_data, masks, and scaler.'),
]
for title_text, desc in steps:
    p = doc.add_paragraph()
    run = p.add_run(title_text + ': ')
    run.bold = True
    p.add_run(desc)

# 3. Graph Construction
doc.add_heading('3. Graph Construction \u2014 gnn_data_prep.py: main()', level=2)
doc.add_paragraph(
    'Called second inside run_training_and_evaluation(). Builds the spatial graph '
    'structure (adjacency matrix) and static node features.'
)

steps2 = [
    ('3a. Load & Concatenate Metadata',
     'load_and_concatenate_metadata() reads coordinate CSVs for all node types: '
     'piezometers (x, y, filter midpoint z), pumping wells, precipitation stations, '
     'evaporation stations, river points. Returns concatenated coordinate arrays, '
     'node type indicators, and counts (num_piezo=200, num_pump=4, num_prec=2, '
     'num_evap=2, num_river=~4). Total ~212 nodes.'),
    ('3b. Create Static Features',
     'create_static_features() normalizes x, y, z coordinates to [0,1] and combines '
     'with node type indicator. Returns a [212 x 4] tensor.'),
    ('3c. Generate Adjacency Matrix',
     'Dispatches to one of several graph generation functions based on graph_type: '
     '"default" (distance-based, fixed weights), "geolayer" (same-layer connectivity '
     '+ distance + optional RF perturbation), "rf" fixed (RF importance selects neighbors, '
     'fixed weight 0.1), "rf" variable (RF importance determines both neighbors and edge '
     'weights). All variants connect: Piezo-Piezo (top-k nearest), Piezo-Pump (closest n), '
     'Piezo-Precip (closest 1), Piezo-Evap (closest 1), Piezo-River (closest 2). '
     'Returns symmetric A_tilde [212 x 212].'),
]
for title_text, desc in steps2:
    p = doc.add_paragraph()
    run = p.add_run(title_text + ': ')
    run.bold = True
    p.add_run(desc)

# 4. Model Creation
doc.add_heading('4. Model Creation', level=2)
doc.add_paragraph(
    'create_mtgnn_model() initializes either MTGNN or MTGNN_LSTM based on '
    'config["model_type"]. The MTGNN architecture consists of:'
)
items = [
    'Start convolution: maps input [B, 1, N, T] to [B, residual_channels, N, T]',
    '4 stacked MTGNNLayers, each containing:',
    '  - DilatedInception: temporal convolution with multiple kernel sizes [1, 2]',
    '  - MixProp: multi-hop graph convolution on A_tilde (spatial mixing)',
    '  - LayerNorm + residual connection + skip connection',
    '  - Dropout (0.5)',
    'End convolution: aggregates skip connections to [B, 1, N, 1]',
    'Output is sliced to piezo nodes only: [B, 1, 200, 1]',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

# 5. Training
doc.add_heading('5. Training \u2014 train()', level=2)
doc.add_paragraph(
    'Training uses progressive future windows and autoregressive prediction.'
)

p = doc.add_paragraph()
run = p.add_run('Progressive Windows: ')
run.bold = True
p.add_run(
    'The model is trained for F_w = 1, then 2, then 3 future steps. Each window '
    'warm-starts from the best model of the previous window.'
)

p = doc.add_paragraph()
run = p.add_run('AutoregressiveTimeSeriesDataset: ')
run.bold = True
p.add_run(
    'Each sample provides: '
    'input = data[idx : idx+W, :200] (W timesteps of piezometer data), '
    'external_forces = data[idx : idx+W+F, 200:] (extended exogenous window), '
    'target = data[idx+W : idx+W+F, :200] (future piezometer values), '
    'mask = missing data mask for the target window.'
)

p = doc.add_paragraph()
run = p.add_run('Autoregressive Loop: ')
run.bold = True
p.add_run(
    'For each future step t, the model predicts one step ahead. The prediction is '
    'appended to the input window (oldest timestep dropped) to predict the next step. '
    'External forces are sliced to provide the correct temporal context at each step.'
)

p = doc.add_paragraph()
run = p.add_run('Loss & Optimization: ')
run.bold = True
p.add_run(
    'Masked MSE loss (predictions * mask vs targets * mask) ensures missing data does '
    'not contribute to gradients. Uses Adam optimizer with ReduceLROnPlateau scheduler '
    'and early stopping.'
)

# 6. Evaluation
doc.add_heading('6. Evaluation', level=2)
doc.add_paragraph(
    'After training, the model is evaluated on the test set:'
)
items2 = [
    'make_predictions(): 100-step autoregressive rollout on a test sample',
    'inverse_transform_with_shape_adjustment(): undo MinMaxScaler normalization',
    'calculate_rmse_per_piezometer(): RMSE per well, summarized by geolayer',
    'Visualization: comparison sequences, dual-Y plots, 3D RMSE network plot',
    'Results saved to training_results/ and summary CSV',
]
for item in items2:
    doc.add_paragraph(item, style='List Bullet')

# 7. Data Shapes
doc.add_heading('Key Data Shapes at Each Stage', level=1)

table = doc.add_table(rows=8, cols=2)
table.style = 'Light Shading Accent 1'
table.rows[0].cells[0].text = 'Stage'
table.rows[0].cells[1].text = 'Shape'

data_rows = [
    ('Raw piezometer CSVs', '~375 files, irregular timestamps'),
    ('After selection', '200 piezometers x T weekly timesteps'),
    ('Combined DataFrame', '212 nodes x T (200 piezo + 4 pump + 2 prec + 2 evap + 4 river)'),
    ('Adjacency matrix', '[212 x 212] sparse, symmetric'),
    ('Static features', '[212 x 4] (normalized x, y, z, type)'),
    ('Model input per batch', '[B, 1, 212, W+1]'),
    ('Model output per step', '[B, 1, 200, 1] (predicts piezo only)'),
]
for i, (stage, shape) in enumerate(data_rows):
    table.rows[i+1].cells[0].text = stage
    table.rows[i+1].cells[1].text = shape

doc_path = out_dir / 'Program_Flow_Documentation.docx'
doc.save(str(doc_path))
print(f"Word document saved: {doc_path}")
